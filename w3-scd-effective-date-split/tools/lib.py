# Shared styling/layout helpers. Mirrors Final_Scenarios_v1.xlsx:
# sources anchored at column A, target block to the right of them.
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

INK = "14213D"; TEAL = "2A6F6B"; RED_T = "FF0000"
HDR   = PatternFill("solid", fgColor="14213D")
GREEN = PatternFill("solid", fgColor="D9EAD3")   # target / result blocks
RED   = PatternFill("solid", fgColor="F7E2DA")   # a source that changed
GREY  = PatternFill("solid", fgColor="F2F2F2")   # a source that did not
AMBER = PatternFill("solid", fgColor="FCF3E8")   # open / undecided
thin  = Side(style="thin", color="BFBFBF")
BOX   = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP  = Alignment(wrap_text=True, vertical="top")
HIGH  = "9999-12-31"


def H(*vals):
    """Readable stand-in for the hash of the target-bound columns."""
    return "H(" + "|".join("" if v in (None, "", "(blank)") else v for v in vals) + ")"


class Sheet:
    def __init__(self, wb, name, title, tabcolor=None):
        self.ws = wb.create_sheet(name)
        self.ws.sheet_view.showGridLines = False
        if tabcolor:
            self.ws.sheet_properties.tabColor = tabcolor
        c = self.ws.cell(row=1, column=1, value=title)
        c.font = Font(bold=True, size=13, color=INK)
        self.wneed = {}

    def _w(self, col, text):
        n = len(str(text)) if text is not None else 0
        self.wneed[col] = max(self.wneed.get(col, 10), min(n + 3, 56))

    def label(self, row, col, text, color=TEAL):
        c = self.ws.cell(row=row, column=col, value=text)
        c.font = Font(bold=True, size=11, color=color)
        if len(str(text)) < 26:
            self._w(col, text)

    def block(self, row, col, heading, cols, rows, fill=None, heading_color=TEAL):
        """Writes a labelled table. Returns the row after the last data row."""
        if heading:
            self.label(row, col, heading, heading_color)
            row += 1
        for i, v in enumerate(cols):
            c = self.ws.cell(row=row, column=col + i, value=v)
            c.font = Font(bold=True, size=10, color="FFFFFF")
            c.fill = HDR; c.border = BOX
            c.alignment = Alignment(wrap_text=True, vertical="center")
            self._w(col + i, v)
        self.ws.row_dimensions[row].height = 26
        row += 1
        for r in rows:
            for i, v in enumerate(r):
                c = self.ws.cell(row=row, column=col + i, value=v)
                c.font = Font(size=10, color=INK); c.border = BOX; c.alignment = WRAP
                if fill:
                    c.fill = fill
                self._w(col + i, v)
            row += 1
        return row

    def _para(self, row, items, heading, hcolor, marker, span):
        if heading:
            c = self.ws.cell(row=row, column=1, value=heading)
            c.font = Font(bold=True, size=11, color=hcolor)
            row += 1
        for t in items:
            c = self.ws.cell(row=row, column=1, value=marker + t)
            c.font = Font(size=10, color=hcolor if marker == "?  " else INK)
            c.alignment = WRAP
            self.ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
            self.ws.row_dimensions[row].height = max(15, 14 * (1 + len(t) // 118))
            row += 1
        return row

    def bullets(self, row, items, heading="How this works", span=8):
        return self._para(row, items, heading, INK, "•  ", span)

    def questions(self, row, items, heading="Open question", span=8):
        return self._para(row, items, heading, RED_T, "?  ", span)

    def stack(self, col, blocks, start=3):
        """Lays blocks down one column, each separated by a blank row.
        blocks: list of (heading, cols, rows, fill). Returns next free row."""
        r = start
        for heading, cols, rows, fill in blocks:
            r = self.block(r, col, heading, cols, rows, fill) + 1
        return r

    def finish(self):
        for col, w in self.wneed.items():
            self.ws.column_dimensions[get_column_letter(col)].width = w
        self.ws.freeze_panes = "A2"
